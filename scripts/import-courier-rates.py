"""Read supplier workbook into private local pricing data. Never exports to public assets."""
import json, sys, hashlib
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
root = Path(__file__).resolve().parent.parent
book = openpyxl.load_workbook(source, data_only=True)
services = []

def add(identifier, name, sheet, rows, weight_col, zone_cols, country_cols, direction='export'):
    ws = book[sheet]
    zones = {}
    for country_col, zone_col in country_cols:
        for row in ws:
            country, zone = row[country_col-1].value, row[zone_col-1].value
            if isinstance(country, str) and str(zone) in zone_cols:
                # Postcode-specific mappings are deliberately excluded from generic country quotes.
                if any(t in country.lower() for t in ['postal code', 'post code', 'postcode']):
                    continue
                zones[country.strip().rstrip('*')] = str(zone)
    rates = []
    for i in rows:
        weight = ws.cell(i, weight_col).value
        if not isinstance(weight, (int, float)):
            continue
        prices = {zone: ws.cell(i, col).value for zone, col in zone_cols.items()}
        if not all(isinstance(v, (int, float)) and v > 0 for v in prices.values()):
            raise ValueError(f'Missing cached price in {sheet}, row {i}')
        rates.append({'weight': weight, 'prices': prices})
    if not rates or not zones:
        raise ValueError(f'No usable rates or zones in {sheet}')
    services.append({'id': identifier, 'name': name, 'direction': direction,
                     'sourceSheet': sheet, 'zones': zones, 'rates': rates})

add('fedex-export', 'FedEx International Priority', 'FedEx Export', range(17,158), 1,
    {chr(65+i): 2+i for i in range(17)}, [(21,24),(27,30),(33,36)])
add('fedex-import', 'FedEx International Priority Import', 'FedEx Import', range(17,157), 1,
    {chr(65+i): 2+i for i in range(12)}, [(16,19),(22,25),(28,31)], 'import')
add('ups-export', 'UPS Express Saver', 'UPS EXPORT', range(17,57), 3,
    {str(i): 3+i for i in range(1,10)}, [(14,18)])
add('dhl-export', 'DHL Express Worldwide', 'DHL EXPORT', range(14,64), 1,
    {str(i): 2+i for i in range(1,15)}, [(18,20)])
add('aramex-export', 'Aramex International', 'Aramex', range(5,45), 1,
    {str(i): 1+i for i in range(1,14)}, [(16,17),(18,19),(20,21),(22,23)])
data = {'schema':1, 'supplier':'Orangestar', 'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),
        'markup':0.30, 'surchargesIncluded':True, 'confirmedValidityYear':2027,
        'validUntil':None, 'currency':'INR', 'dimensionalDivisor':5000, 'services':services}
target = root / '.local' / 'courier-rates.json'
target.parent.mkdir(exist_ok=True)
target.write_text(json.dumps(data, ensure_ascii=False, separators=(',',':')), encoding='utf-8')
print(json.dumps({'services':[{'name':s['name'],'countries':len(s['zones']),'slabs':len(s['rates']),
                             'maxKg':s['rates'][-1]['weight']} for s in services]}))
